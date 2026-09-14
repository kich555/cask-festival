#!/usr/bin/env python3
"""엑셀(캐스크카니발_참가업체_목록.xlsx) → 참가 업체 데이터 + 로고 폴더 동기화.

- '인스타 업로드 여부'가 O 인 업체만 홈페이지 데이터(brands2026.json)에 넣는다.
- 엑셀의 모든 업체에 대해 public/brands/<한글명>/ 폴더를 만든다.
- 엑셀에서 이름이 바뀌어 짝이 없어진 폴더는 삭제하지 않고 알려만 준다.

사용법:  python3 scripts/import-brands.py
이어서:  node scripts/link-logos.mjs   (폴더 안 로고를 데이터에 연결)
"""
import json
import os
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
EXCEL = ROOT.parent / "캐스크카니발_참가업체_목록.xlsx"
JSON_PATH = ROOT / "src" / "content" / "brands2026.json"
LOGO_DIR = ROOT / "public" / "brands"

COUNTRY = {
    "한국": "KR", "일본": "JP", "대만": "TW",
    "영국": "GB", "스코틀랜드": "GB-SCT", "네덜란드": "NL",
}

def slugify(name: str, used: set) -> str:
    s = unicodedata.normalize("NFKD", name)
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-").lower() or "brand"
    base, n = s, 2
    while s in used:
        s = f"{base}-{n}"
        n += 1
    used.add(s)
    return s

def folder_name(name_ko: str) -> str:
    """폴더명에 쓸 수 없는 문자만 치환한다."""
    return re.sub(r'[/:*?"<>|]', "_", name_ko).strip()

def main() -> None:
    if not EXCEL.exists():
        raise SystemExit(f"엑셀을 찾을 수 없습니다: {EXCEL}")

    ws = load_workbook(EXCEL, data_only=True)["참가업체"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]

    used: set = set()
    brands = []
    for r in rows:
        if (r[7] or "").strip().upper() != "O":
            continue
        name_ko = str(r[1]).strip()
        name_en = str(r[2] or "").strip().replace("\xa0", "")
        country_ko = str(r[3] or "").strip()
        brands.append({
            "slug": slugify(name_en or name_ko, used),
            "nameKo": name_ko,
            "nameEn": name_en,
            "country": COUNTRY.get(country_ko, ""),
            "countryKo": country_ko,
            "website": str(r[4]).strip() if r[4] else None,
            "instagram": str(r[5]).strip() if r[5] else None,
            "booths": int(r[6]) if r[6] else 1,
            "logo": None,
        })

    JSON_PATH.write_text(
        json.dumps(brands, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    # 폴더 동기화 (한글 자모 분리 차이를 없애기 위해 NFC 로 비교)
    LOGO_DIR.mkdir(parents=True, exist_ok=True)
    wanted = {unicodedata.normalize("NFC", folder_name(str(r[1]).strip())) for r in rows}
    existing = {
        unicodedata.normalize("NFC", p.name): p
        for p in LOGO_DIR.iterdir()
        if p.is_dir()
    }

    created = 0
    for name in sorted(wanted):
        if name not in existing:
            (LOGO_DIR / name).mkdir()
            created += 1

    stale = []
    for name, path in existing.items():
        if name in wanted:
            continue
        files = [f for f in path.iterdir() if not f.name.startswith(".")]
        stale.append((name, len(files)))

    print(f"홈페이지 노출 업체 {len(brands)}곳 → {JSON_PATH.relative_to(ROOT)}")
    print(f"업체 폴더: 신규 {created}개 / 전체 {len(wanted)}개")

    if stale:
        print("\n엑셀에 없는 폴더 (이름이 바뀌었을 수 있습니다. 직접 확인해 주세요)")
        for name, count in stale:
            mark = f"파일 {count}개 있음" if count else "비어 있음"
            print(f"  · {name}  — {mark}")

if __name__ == "__main__":
    main()
